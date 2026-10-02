"""부적합중량 산출용 단위중량표 로딩.
- 재공품/조립: '재공품 중량정보.xls' (재공품코드별 Wet/Dry중량(g)) — 재공품은 Wet, 조립은 Dry 기준(확정)
- 완성: '제품 중량정보.xlsx' (자재코드별 총중량(kg) → g로 환산)
매년/수시로 갱신될 수 있어 config에 기본 번들하되 화면에서 교체 가능해야 한다.

(v22.0) 사용자가 "제품중량, 제조원가, 극판중량_260922.xlsx"(SAP 자재마스터 원본, 시트 3개짜리 통합
파일)를 제공 — 극판중량/제품중량 둘 다 이 한 파일의 시트로도 올 수 있어 시트 이름으로 자동 탐색하는
_find_sheet()를 추가(cost.load_cost_lookup()과 동일 패턴). 제품중량 시트는 컬럼명이 기존 번들 파일과
달리 "총중량(kg)" 하나가 아니라 "총중량"+"중량단위"로 분리돼 있어 그 형식도 함께 지원한다."""
import pandas as pd


def _find_sheet(path_or_file, keyword: str) -> str:
    """워크북 시트 이름 중 keyword가 포함된 것을 찾는다(없으면 첫 시트) — cost.py와 동일 패턴."""
    import openpyxl
    wb = openpyxl.load_workbook(path_or_file, read_only=True, data_only=True)
    for name in wb.sheetnames:
        if keyword in name:
            return name
    return wb.sheetnames[0]


def load_plate_weight_lookup(path_or_file) -> dict:
    """{재공품코드: {"wet": g, "dry": g}} 반환."""
    try:
        sheet = _find_sheet(path_or_file, "극판중량")
        df = pd.read_excel(path_or_file, sheet_name=sheet)
    except Exception:
        return {}
    df = df.rename(columns=lambda c: str(c).strip())
    required = {"재공품코드", "Wet중량(g)", "Dry중량(g)"}
    if not required.issubset(df.columns):
        return {}
    df = df.dropna(subset=["재공품코드"])
    df = df.drop_duplicates(subset="재공품코드", keep="last")
    return {
        row["재공품코드"]: {"wet": row["Wet중량(g)"], "dry": row["Dry중량(g)"]}
        for _, row in df.iterrows()
    }


def load_product_weight_lookup(path_or_file) -> dict:
    """{자재코드: 총중량(g)} 반환 (원본은 kg 단위 → g로 환산)."""
    try:
        sheet = _find_sheet(path_or_file, "제품중량")
        df = pd.read_excel(path_or_file, sheet_name=sheet)
    except Exception:
        return {}
    df = df.rename(columns=lambda c: str(c).strip())
    if "자재코드" not in df.columns:
        return {}
    if "총중량(kg)" in df.columns:
        weight_kg = df["총중량(kg)"]
    elif {"총중량", "중량단위"}.issubset(df.columns):
        # SAP 자재마스터 원본 형식 — 단위가 "KG"로 명시된 행만 사용(그 외 단위/결측은 모름 처리,
        # 임의 환산하지 않음 — CLAUDE.md 데이터 정확성 원칙).
        is_kg = df["중량단위"].astype(str).str.strip().str.upper() == "KG"
        weight_kg = df["총중량"].where(is_kg)
    else:
        return {}
    df = df.assign(_총중량_kg=weight_kg).dropna(subset=["자재코드", "_총중량_kg"])
    df = df.drop_duplicates(subset="자재코드", keep="last")
    return dict(zip(df["자재코드"], df["_총중량_kg"] * 1000))


def compute_defect_weight_g(process: str, code, qty: float,
                             plate_weight_lookup: dict, product_weight_lookup: dict):
    """공정별 확정 기준(재공품=Wet, 조립=Dry, 완성=제품 총중량)으로 불량수량×단위중량을 계산한다.
    단가표에 코드가 없으면 None(모름) — 0으로 임의 대체하지 않고 호출부에서 표시만 0으로 채운다."""
    unit = None
    if process == "재공품":
        unit = plate_weight_lookup.get(code, {}).get("wet")
    elif process == "조립":
        unit = plate_weight_lookup.get(code, {}).get("dry")
    elif process == "완성":
        unit = product_weight_lookup.get(code)
    if unit is None:
        return None
    return (qty or 0) * unit
