"""제조원가 단가표('25년 제조원가, 단가.xlsx' 등) 로딩.
자재코드(=매칭키: 품명/재공품코드/제품코드)별 표준원가(원/개)를 조회할 수 있게 한다.
매년 파일이 바뀌므로 기본값으로 번들해두되, 사이드바에서 새 파일로 교체 가능해야 함(config.COST_XLSX 참고)."""
import openpyxl
import pandas as pd


def _find_sheet(path_or_file, keyword: str) -> str:
    wb = openpyxl.load_workbook(path_or_file, read_only=True, data_only=True)
    for name in wb.sheetnames:
        if keyword in name:
            return name
    return wb.sheetnames[0]


def load_cost_lookup(path_or_file) -> dict:
    """{자재코드: 표준원가(원/개)} 반환. 같은 코드가 여러 (년도,기간)에 있으면 최신 것을 쓴다."""
    try:
        sheet = _find_sheet(path_or_file, "제조원가")
        df = pd.read_excel(path_or_file, sheet_name=sheet)
    except Exception:
        return {}

    df = df.rename(columns=lambda c: str(c).strip())
    required = {"년도", "기간", "자재", "표준원가"}
    if not required.issubset(df.columns):
        return {}

    df["년도"] = pd.to_numeric(df["년도"], errors="coerce")
    df["기간"] = pd.to_numeric(df["기간"], errors="coerce")
    df = df.dropna(subset=["자재", "표준원가"])
    df = df.sort_values(["년도", "기간"], ascending=False)
    df = df.drop_duplicates(subset="자재", keep="first")
    return dict(zip(df["자재"], df["표준원가"]))
